"""导出队列与产物清理（P5-3 / F5-2~F5-6 / P5-A5 / P5-A6 / P5-C1）。

这一份测的是**功能**，不是「函数返回了什么」：造一门真课、真排版、真落盘，
然后检查盘上那个文件、库里那一行、以及一份记录读起来对不对。

外部能力一个都不用：三种渲染器都是纯函数（python-pptx / PyMuPDF 都在本地跑），
不调模型、不连网。所以这一份在任何机器上都该全绿。

**导出目录必须落在临时目录里**（`conftest.base_env` 的 `EXPORT_DIR`）。
否则测试会往 `backend/data/exports/` 里写 —— 那是开发机上用户真正的产物目录。
最后一个用例专门守着这条线。
"""

from __future__ import annotations

from datetime import timedelta

import pytest

from app.common.errors import NotFoundError, StateError, ValidationError
from app.common.timeutil import parse_iso, utcnow

TOPIC = "机器学习入门"


@pytest.fixture()
def exported(app, tmp_path):
    """一门三页的课 + 一节上完的课堂记录，导出目录指到 `tmp_path`。

    导出目录在**建 app 之后**再改一次：`base_env` 已经把它放到库文件旁边了，
    这里挪到 `tmp_path` 是为了每个用例拿到自己那份干净目录（`tmp_path` 是函数级的）。
    """
    from app.extensions import db
    from app.models import (
        AgentRole,
        ClassroomMessage,
        ClassroomSession,
        Course,
        CoursePage,
        Material,
        MaterialChunk,
        PageSource,
        QuizAttempt,
        User,
    )

    app.config["EXPORT_DIR"] = str(tmp_path / "exports")

    # 表之间没有 ORM 关系，SQLAlchemy 推不出插入顺序，所以按依赖分三次 flush
    db.session.add(User(id="u_teacher", name="沈老师", role="teacher"))
    db.session.add(User(id="u_stu", name="小周", role="student"))
    db.session.add(AgentRole(code="teacher", name="沈老师", role="teacher", builtin=True))
    db.session.add(AgentRole(code="s1", name="小林", role="student", builtin=True))
    db.session.flush()

    course = Course(id="c_1", owner_id="u_teacher", title=TOPIC, topic=TOPIC,
                    status="ready", page_count=3)
    # DSL 的键名由 `ir._page_from_dsl` / `ir._blocks_for` 定：内容按**页型**
    # 从页面自己的字段里取（`bullets` / `quiz` / `narration`），不是一个通用的
    # `blocks` 数组。这里照着写，导出的产物里才真有东西。
    course.dsl = {
        "meta": {"title": TOPIC, "subtitle": "从零开始", "topic": TOPIC,
                 "audience": "大一新生", "durationMin": 20,
                 "cover": {"lecturer": "沈老师", "durationMin": 20, "audience": "大一新生"}},
        "chapters": [{"no": 1, "title": "什么是机器学习", "summary": "", "pages": [1, 2, 3]}],
        "pages": [
            {"pageNo": 1, "chapterNo": 1, "kind": "cover", "title": TOPIC,
             "subtitle": "从零开始", "meta": {"lecturer": "沈老师", "durationMin": 20}},
            {"pageNo": 2, "chapterNo": 1, "kind": "concept", "title": "什么是机器学习",
             "narration": [{"beatId": "b1", "text": "它让机器从数据里找规律。"}],
             "bullets": [{"text": "从数据里找规律", "emphasis": ["数据"]},
                         {"text": "不用人写规则"}]},
            {"pageNo": 3, "chapterNo": 1, "kind": "quiz", "title": "随堂小测",
             "quiz": {"stem": "机器学习靠什么？", "options": ["数据", "运气"],
                      "answer": "数据", "explain": "从数据里找规律。", "tag": "选择题"}},
        ],
    }
    db.session.add(course)
    db.session.flush()

    for no in (1, 2, 3):
        page = CoursePage(id=f"p{no}", course_id="c_1", page_no=no, chapter_no=1,
                          kind="concept", title=f"第 {no} 页", status="ready")
        page.dsl = {}
        db.session.add(page)

    db.session.add(Material(id="m1", owner_id="u_teacher", name="机器学习讲义.md",
                            ext="md", size_bytes=10, status="ready"))
    db.session.flush()
    db.session.add(MaterialChunk(id="k1", material_id="m1", chunk_no=1,
                                 page_from=3, page_to=3, text="机器学习是……"))
    db.session.add(PageSource(id="ps1", page_id="p2", material_id="m1", chunk_id="k1",
                              page_no=3, section_path="1.1", quote="机器学习是……", score=0.9))

    db.session.add(ClassroomSession(
        id="s_1", course_id="c_1", owner_id="u_teacher", mode="auto", status="ended",
        started_at="2026-09-13T10:00:00.000000", ended_at="2026-09-13T10:42:00.000000",
    ))
    db.session.add(ClassroomMessage(
        id="msg1", session_id="s_1", speaker_code="teacher", speaker_kind="teacher",
        type="lecture", text="大家好，今天我们讲机器学习。",
        page_no=1, beat_id="b1", ts="2026-09-13T10:00:05.000000",
    ))
    db.session.add(ClassroomMessage(
        id="msg2", session_id="s_1", speaker_code="u_stu", speaker_kind="me",
        type="question", text="老师，数据和样本是一回事吗？",
        page_no=2, ts="2026-09-13T10:12:00.000000",
    ))
    db.session.add(QuizAttempt(
        id="q1", session_id="s_1", course_id="c_1", page_no=2, option="数据",
        correct=True, response_ms=4200, ts="2026-09-13T10:20:00.000000",
    ))
    db.session.commit()
    return course


def run_export(course, **kwargs):
    """建一次导出并同步跑完，返回库里那一行的最新状态。"""
    from app.extensions import db
    from app.models import Export
    from app.services.exports import queue

    row = queue.create(course, **kwargs)
    queue.run(row.id)
    return db.session.get(Export, row.id)


def _row_whose_session_vanished(course):
    """建一次课堂记录导出，再把那堂课删掉 —— 轮到它渲染时记录已经重建不出来。

    真实世界里这一步是「排队期间有人删了那堂课」：`create()` 拦不住它，
    它只在建的那一刻判一次归属。而这正是 `run()` 必须把失败写进库的原因 ——
    失败发生时请求早就返回了，没有别人能把这个错误交给用户。
    """
    from app.extensions import db
    from app.models import ClassroomSession
    from app.services.exports import queue

    session = ClassroomSession(course_id=course.id, owner_id="u_teacher", status="lecture")
    db.session.add(session)
    db.session.commit()

    row = queue.create(course, fmt="md", scope="record", session_id=session.id)
    db.session.delete(session)
    db.session.commit()
    return row


# --------------------------------------------------------------------------
# 支持矩阵
# --------------------------------------------------------------------------


def test_supported_formats_by_scope(exported):
    from app.services.exports import queue

    assert queue.supported_formats("course") == ("pptx", "html", "pdf")
    assert queue.supported_formats("record") == ("md", "pdf")
    assert queue.supported_formats("nonsense") == ()


@pytest.mark.parametrize(
    ("scope", "fmt"),
    [("course", "md"), ("record", "pptx"), ("record", "html"), ("course", "docx")],
)
def test_unsupported_combo_rejected_upfront(exported, scope, fmt):
    """不支持的组合在 `create()` 就挡住，不留到渲染时才发现（用户要等 3 秒）。"""
    from app.models import Export
    from app.services.exports import queue

    with pytest.raises(ValidationError) as info:
        queue.create(exported, fmt=fmt, scope=scope, session_id="s_1")
    assert fmt in str(info.value.message)
    assert Export.query.count() == 0, "被拒的请求不该留下行"


def test_record_export_needs_session_id(exported):
    from app.services.exports import queue

    with pytest.raises(ValidationError):
        queue.create(exported, fmt="md", scope="record")


def test_course_without_ready_pages_cannot_export(app, tmp_path):
    """课程还没生成完就导出 → 40902（`errors.py` 里 StateError 就是照这个场景写的）。"""
    from app.extensions import db
    from app.models import Course, User
    from app.services.exports import queue

    app.config["EXPORT_DIR"] = str(tmp_path / "exports")
    db.session.add(User(id="u_owner", name="空课老师", role="teacher"))
    db.session.flush()
    db.session.add(Course(id="c_empty", owner_id="u_owner", title="空课", topic="空课"))
    db.session.commit()

    with pytest.raises(StateError):
        queue.create(db.session.get(Course, "c_empty"), fmt="pptx")


def test_a_course_whose_dsl_has_no_pages_cannot_export(exported):
    """页面行满着、`courses.dsl` 里的 `pages` 却是空的 → 40902，**不能**导出一份空产物。

    这是两份数据对不上的那一种：`_require_ready` 只看 `course_pages` 会放行，
    而渲染读的是 `courses.dsl_json`。放过去的后果按格式各不相同 ——
    PPTX 与 HTML 安静地给出一个没内容的文件（用户下载完才发现），
    PDF 在 `DocumentWriter.close()` 上抛 `cannot save with zero pages`，
    变成一句「导出失败，请重试」。
    """
    from app.extensions import db
    from app.services.exports import queue

    pages = list((exported.dsl or {}).get("pages") or [])
    assert pages, "布景里这门课本来就该有页面"
    exported.dsl = {**(exported.dsl or {}), "pages": []}
    db.session.commit()
    try:
        with pytest.raises(StateError) as info:
            queue.create(exported, fmt="pdf")
        assert "空" in str(info.value.message)
    finally:
        # 课程对象是夹具给的那一个，改坏了后面的用例会读到一份没有页的 DSL
        exported.dsl = {**(exported.dsl or {}), "pages": pages}
        db.session.commit()


def test_a_record_export_of_a_session_that_does_not_exist_is_refused(exported):
    """课堂记录导出在**建的时候**就判归属，不留到渲染时 ——
    一次越权请求不该先排上队、再变成一个失败行，那是把越权写进了导出历史。"""
    from app.models import Export
    from app.services.exports import queue

    with pytest.raises(NotFoundError):
        queue.create(exported, fmt="md", scope="record", session_id="s_missing")
    assert Export.query.count() == 0


# --------------------------------------------------------------------------
# 课件三种格式
# --------------------------------------------------------------------------


@pytest.mark.parametrize("fmt", ["pptx", "html", "pdf"])
def test_course_export_writes_real_file(exported, fmt):
    from app.services.exports import store

    row = run_export(exported, fmt=fmt, owner_id="u_teacher")

    assert row.status == "done", row.error
    assert row.progress == 100
    assert row.error == ""
    assert row.file_path.startswith(store.SUBDIR + "/"), row.file_path
    assert row.file_path.endswith("." + fmt)

    path = store.file_of(row)
    assert path is not None and path.is_file()
    assert path.stat().st_size == row.size_bytes > 0
    assert row.expires_at, "到期时间是给清理任务的承诺，不能不写"
    assert row.finished_at


def test_pptx_is_editable_slides_not_an_image(exported):
    """F5-2：产物要能被 PowerPoint 打开、每页一张幻灯片、讲稿在**备注页**里。"""
    from pptx import Presentation

    from app.services.exports import store

    row = run_export(exported, fmt="pptx")
    prs = Presentation(str(store.file_of(row)))

    assert len(prs.slides) == 3
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides if s.has_notes_slide]
    assert any("机器从数据里找规律" in text for text in notes), f"讲稿没进备注页：{notes}"
    # 正文要是真的文本框（可编辑），不是一张图
    texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
    assert any("不用人写规则" in text for text in texts)


def test_html_is_self_contained(exported):
    """F5-3：断网也要能打开 —— 一个外链都不能有。"""
    from app.services.exports import store

    row = run_export(exported, fmt="html")
    text = store.file_of(row).read_text(encoding="utf-8")

    assert "http://" not in text and "https://" not in text
    assert "<script" in text and "slide__kind" in text
    assert "机器学习靠什么？" in text and "从数据里找规律" in text


def test_pdf_has_all_pages_and_page_numbers(exported):
    import pymupdf

    from app.services.exports import store

    row = run_export(exported, fmt="pdf")
    doc = pymupdf.open(str(store.file_of(row)))
    try:
        assert doc.page_count >= 3
        text = "".join(doc[i].get_text() for i in range(doc.page_count))
        assert "机器学习" in text, "PDF 里抽不到正文（中文没画出来）"
        assert f"第 {doc.page_count} / {doc.page_count} 页" in doc[doc.page_count - 1].get_text()
    finally:
        doc.close()


def test_watermark_can_be_turned_off_and_is_recorded(exported):
    """水印开着是合规项，关掉也得留档 —— 看到一份没水印的产物要能说清它从哪来。"""
    from app.services.exports import store

    on = run_export(exported, fmt="html")
    off = run_export(exported, fmt="html", options={"watermark": False})

    assert off.options["watermark"] is False
    assert on.options["watermark"] is True

    # 水印那一句也写在模板的 JS 里（`wm.textContent = ...`），所以不能直接搜字符串 ——
    # 要看的是**打给页面的那份开关**，它决定 JS 到底挂不挂那个 div。
    on_text = store.file_of(on).read_text(encoding="utf-8")
    off_text = store.file_of(off).read_text(encoding="utf-8")
    assert '"watermark": true' in on_text
    assert '"watermark": false' in off_text


def test_source_label_is_the_human_readable_one(exported):
    """DSL 里只有 chunkId，人要看的是「《机器学习讲义.md》第 3 页」（P4 的出处）。"""
    from app.services.exports import queue as export_queue

    row = run_export(exported, fmt="pdf")
    page = export_queue._deck_of(row).page(2)

    assert page is not None
    assert [item.label for item in page.sources] == ["《机器学习讲义.md》第 3 页"]
    assert page.sources[0].quote == "机器学习是……"


def test_source_label_survives_a_deleted_material(exported):
    """材料被删了也要给一行能读的 label，而不是 `《》第 3 页` 这种空壳。"""
    from app.services.exports import queue as export_queue

    row = run_export(exported, fmt="pdf")
    assert export_queue._label("", 3) == "材料已删除（原第 3 页）"
    assert export_queue._label("讲义.md", 0) == "《讲义.md》"
    assert export_queue._label("", 0) == "材料已删除"
    assert row.status == "done"


# --------------------------------------------------------------------------
# 课堂记录
# --------------------------------------------------------------------------


def test_record_markdown_reads_like_a_record(exported):
    """F5-5：一份记录要能直接读 —— 时间、人名、谁说了什么、答对没有。"""
    from app.services.exports import store

    row = run_export(exported, fmt="md", scope="record", session_id="s_1")
    assert row.status == "done", row.error

    text = store.file_of(row).read_text(encoding="utf-8")
    assert f"# {TOPIC} · 课堂记录" in text
    assert "2026-09-13 10:00" in text
    assert "10:42" in text
    assert "**主讲**：沈老师" in text
    assert "用时 42 分钟" in text
    assert "大家好，今天我们讲机器学习。" in text
    assert "答对" in text
    # 栏目标题比文档标题低一级，否则读的人分不清哪句是标题哪句是栏目
    assert "\n## 课堂概览\n" in text
    assert "\n## 字幕全文\n" in text


def test_record_resolves_speaker_codes_to_names(exported):
    """`speaker_code` 对 AI 是 `agent_roles.code`、对真人是 `user_id` —— 两个都要还原。

    交出去一份写着 `speaker: s1` 的记录，读的人根本不知道那是谁。
    """
    from app.services.exports import store

    row = run_export(exported, fmt="md", scope="record", session_id="s_1")
    text = store.file_of(row).read_text(encoding="utf-8")

    assert "沈老师" in text, "AI 角色的名字没还原"
    assert "小周" in text, "真人的名字没还原"
    assert "u_teacher" not in text and "u_stu" not in text, "id 泄进了正文"


def test_record_keeps_raw_code_when_nobody_claims_it(exported):
    """两边都查不到时**原样显示 code**：显示成「未知」等于把「这是谁」彻底丢掉。"""
    from app.services.exports.record import _speaker

    assert _speaker("s9", {}) == "s9"
    assert _speaker("", {}) == "（未署名）"
    assert _speaker("teacher", {"teacher": "沈老师"}) == "沈老师"


def test_record_separates_subtitles_from_discussion(exported):
    """字幕（带 beatId）与讨论区（不带）分成两栏，同一条发言不会两边都算。"""
    from app.services.exports import store

    row = run_export(exported, fmt="md", scope="record", session_id="s_1")
    text = store.file_of(row).read_text(encoding="utf-8")

    subtitles, discussion = text.split("## 讨论区")
    assert "大家好，今天我们讲机器学习。" in subtitles
    assert "大家好，今天我们讲机器学习。" not in discussion
    assert "数据和样本是一回事吗" in discussion


def test_record_pdf_reuses_the_pdf_pipeline(exported):
    """记录也出 PDF，走的是同一套排版与页脚（不是第二份实现）。"""
    import pymupdf

    from app.services.exports import store

    row = run_export(exported, fmt="pdf", scope="record", session_id="s_1")
    assert row.status == "done", row.error

    doc = pymupdf.open(str(store.file_of(row)))
    try:
        text = "".join(doc[i].get_text() for i in range(doc.page_count))
        assert "课堂记录" in text and "沈老师" in text
        assert f"第 {doc.page_count} / {doc.page_count} 页" in doc[doc.page_count - 1].get_text()
    finally:
        doc.close()


def test_record_of_a_quiet_session_says_so(exported):
    """空栏目也出一页并写明「没有」—— 「没有」与「漏导了」是两件事。"""
    from app.extensions import db
    from app.models import ClassroomSession
    from app.services.exports import store

    db.session.add(ClassroomSession(
        id="s_empty", course_id="c_1", owner_id="u_teacher", mode="auto", status="idle",
    ))
    db.session.commit()

    row = run_export(exported, fmt="md", scope="record", session_id="s_empty")
    text = store.file_of(row).read_text(encoding="utf-8")

    assert "这堂课没有字幕" in text
    assert "这堂课没有讨论区发言" in text
    assert "尚未下课" in text


# --------------------------------------------------------------------------
# 失败与进度
# --------------------------------------------------------------------------


def test_failure_is_recorded_in_the_row_not_raised(exported):
    """失败必须落到库里 —— 抛出去只有线程池的日志看得到，用户那边是一行卡住的 running。"""
    from app.extensions import db
    from app.models import Export
    from app.services.exports import queue

    row = _row_whose_session_vanished(exported)
    queue.run(row.id)  # 不抛异常 —— 这是重点

    fresh = db.session.get(Export, row.id)
    assert fresh.status == "failed"
    assert fresh.error, "失败必须留下给人看的原因"
    assert fresh.file_path is None and fresh.size_bytes == 0
    assert fresh.finished_at, "失败也要记时刻"
    assert fresh.expires_at == "", "没跑成的产物没有到期承诺，清理时不该碰它"


def test_progress_is_written_while_rendering(exported, monkeypatch):
    """进度条的意义是「现在进行到哪了」，所以渲染途中就要有人写它。

    这里拦住渲染器的**第 1 页回调**去看库里那一行的进度 —— 那一刻渲染还没结束，
    如果进度是攒到最后一起写的，这里读到的就是 5。
    """
    from app.extensions import db
    from app.models import Export
    from app.services.exports import html as html_renderer
    from app.services.exports import queue

    seen: list[int] = []
    original = html_renderer.document

    def spy(deck, options, *, generated_at="", on_page=None):
        def hook(done, total):
            if on_page is not None:
                on_page(done, total)
            if done == 1:
                seen.append(db.session.get(Export, row_id).progress)
        return original(deck, options, generated_at=generated_at, on_page=hook)

    monkeypatch.setattr(html_renderer, "document", spy)
    row = queue.create(exported, fmt="html", owner_id="u_teacher")
    row_id = row.id
    queue.run(row_id)

    assert seen, "第 1 页的回调没被调用"
    assert seen[0] > 5, f"渲染途中进度还停在起点：{seen}"
    assert db.session.get(Export, row_id).progress == 100


def test_background_submit_runs_to_completion(app, exported):
    """`submit()` 走的才是生产路径（请求线程立即返回，渲染在后台线程）。"""
    from app.common.tasks import shutdown_runner, wait_for
    from app.extensions import db
    from app.models import Export
    from app.services.exports import queue

    row = queue.create(exported, fmt="html", owner_id="u_teacher")
    assert queue.submit(row.id) is True
    assert wait_for(queue.KEY_PREFIX + row.id, timeout=60) is True

    # 后台线程用的是**它自己的会话**，这边会话的身份映射里还是那个 `queued` 的对象 ——
    # 不 expire 一下读到的就是旧的（`test_p4_material_api` 里同一条理由）。
    db.session.expire_all()
    done = db.session.get(Export, row.id)
    assert done.status == "done", done.error
    shutdown_runner(app)


def test_export_keys_are_namespaced_away_from_generation_jobs():
    """导出任务的 key 有前缀：某次导出的 id 撞上一个生成任务 id 时不会互相挡住。"""
    from app.services.exports.queue import KEY_PREFIX

    assert KEY_PREFIX.startswith("export:")
    assert KEY_PREFIX + "abc" != "abc"


# --------------------------------------------------------------------------
# 清理（P5-A6 / P5-C1）
# --------------------------------------------------------------------------


def test_cleanup_deletes_expired_files_and_rows(exported):
    from app.models import Export
    from app.services.exports import queue, store

    run_export(exported, fmt="html")
    assert store.used_bytes() > 0

    report = queue.cleanup(now=utcnow() + timedelta(hours=48))

    assert report["exports"] == 1
    assert report["files"] == 1
    assert Export.query.count() == 0
    assert store.used_bytes() == 0


def test_cleanup_leaves_unexpired_and_failed_rows_alone(exported):
    """没过期的不动；失败的行**没有 `expires_at`**，不清 —— 它是「上次为什么失败」的唯一线索。"""
    from app.extensions import db
    from app.models import Export
    from app.services.exports import queue

    ok = run_export(exported, fmt="html")
    bad_row = _row_whose_session_vanished(exported)
    queue.run(bad_row.id)

    queue.cleanup(now=utcnow() + timedelta(hours=1))

    assert db.session.get(Export, ok.id).status == "done"
    assert db.session.get(Export, bad_row.id).status == "failed"
    assert db.session.get(Export, bad_row.id).error


def test_cleanup_sweeps_orphans_only_after_the_grace_period(exported):
    """崩在「写完文件」与「改库」之间的产物，顺着行是找不到的 —— 得扫盘。

    但**宽限期内绝不能动**：那一刻文件可能正是另一次导出刚写下去、
    行还没提交的那一份。
    """
    from app.models import Export
    from app.services.exports import queue, store

    run_export(exported, fmt="html")
    folder = store.physical_path(store.rel_dir("c_1"))
    assert folder is not None and folder.is_dir()
    orphan = folder / "abandoned.html"
    orphan.write_bytes(b"x" * 128)

    assert queue.cleanup(now=utcnow() + timedelta(hours=48))["orphans"] == 0, "刚写下的被误删了"
    assert orphan.is_file()
    assert store.used_bytes() == orphan.stat().st_size  # 过期那份走了，只剩孤儿

    alive = [item.file_path for item in Export.query.all() if item.file_path]
    assert store.sweep_orphans(alive) == 0, "宽限期内不该动手"
    assert store.sweep_orphans(alive, max_age=0) == 1
    assert not orphan.exists()


def test_orphan_sweep_never_touches_an_owned_product(exported):
    """有主的产物哪怕再旧也不能删 —— 这是这个清扫最要紧的一条。"""
    from app.services.exports import store

    row = run_export(exported, fmt="html")
    path = store.file_of(row)
    assert path is not None

    assert store.sweep_orphans([row.file_path], max_age=0) == 0
    assert path.is_file()

    # 反过来：不把它报成「活着」，它就是个孤儿 —— 说明上面那条不是因为没扫到
    assert store.sweep_orphans([], max_age=0) == 1
    assert not path.exists()


def test_temp_files_are_never_taken_for_products(exported):
    """`.export-` 临时文件归 `clean_tmp` 管，孤儿清扫不碰它 —— 它可能正是**正在写**的那一份。"""
    from app.services.exports import store

    row = run_export(exported, fmt="html")
    folder = store.physical_path(store.rel_dir("c_1"))
    assert folder is not None and folder.is_dir()
    tmp = folder / (store.TMP_PREFIX + "abc")
    tmp.write_bytes(b"partial")
    try:
        # 真实产物报成「活着」，于是这一轮唯一够格当孤儿的候选就是那个临时文件
        assert store.sweep_orphans([row.file_path], max_age=0) == 0
        assert tmp.is_file(), "临时文件被当成无主产物删了"
        assert store.file_of(row) is not None
    finally:
        tmp.unlink(missing_ok=True)


def test_purge_course_removes_the_folder_but_not_the_root(exported):
    """删课时端掉 `{EXPORT_DIR}/{course_id}/`，绝不往导出根上删。"""
    from app.services.exports import store

    run_export(exported, fmt="html")
    root = store.root()

    assert store.purge_course("c_1") == 1
    assert not (root / store.rel_dir("c_1")).exists()
    assert root.is_dir(), "导出根目录被删掉了"

    # 不认识的 course_id 是空操作，不是「删根目录」
    assert store.purge_course("") == 0
    assert store.purge_course("../../") == 0
    assert root.is_dir()


def test_expires_at_honours_the_configured_ttl(exported):
    from app.services.exports import store

    row = run_export(exported, fmt="html")

    ttl_hours = store.ttl_hours()
    assert ttl_hours > 0
    delta = parse_iso(row.expires_at) - parse_iso(row.finished_at)
    assert abs(delta.total_seconds() - ttl_hours * 3600) < 120


def test_oversized_product_is_refused_before_writing(exported, monkeypatch):
    """超过上限的产物**根本不写** —— 写了一半再删是白花一次 IO。"""
    from app.extensions import db
    from app.models import Export
    from app.services.exports import queue, store

    row = queue.create(exported, fmt="html", owner_id="u_teacher")
    monkeypatch.setattr(store, "max_bytes", lambda: 16)
    queue.run(row.id)

    fresh = db.session.get(Export, row.id)
    assert fresh.status == "failed"
    assert "过大" in fresh.error
    assert fresh.file_path is None
    assert store.used_bytes() == 0


def test_export_dir_is_not_the_real_data_dir(app):
    """守住测试自己的边界：产物必须落在临时目录，不能写进 backend/data/exports。"""
    from app.config import BACKEND_DIR
    from app.services.exports import store

    resolved = store.root().resolve()
    assert resolved != BACKEND_DIR and BACKEND_DIR not in resolved.parents, (
        f"测试正在往真实的导出目录写：{resolved}"
    )
